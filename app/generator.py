# app/generator.py
# Módulo de generación de documentos ATS-friendly
# ---------------------------------------------------------------------------
# Formatos soportados:
#   1. .docx (Microsoft Word) usando python-docx
#   2. .html (página web ATS-semántica) con template responsive
# ---------------------------------------------------------------------------

import json
import logging
from pathlib import Path
from typing import Optional

from docx import Document as DocumentBuilder
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import CVData, JobTarget

log = logging.getLogger(__name__)

# ============================================================================
# CONSTANTES DE ESTILO
# ============================================================================

FONT_NAME = "Calibri"
FONT_SIZE_BODY = 11
FONT_SIZE_TITLE = 16
FONT_SIZE_SECTION = 13
COLOR_DARK = RGBColor(0x1A, 0x23, 0x2A)


# ============================================================================
# GENERADOR .DOCX
# ============================================================================

def _set_font(
    run,
    name: str = FONT_NAME,
    size: int = FONT_SIZE_BODY,
    bold: bool = False,
    color: Optional[RGBColor] = None,
):
    """Configura fuente en un Run de python-docx."""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def generate_ats_docx(
    cv: CVData,
    output_path: str,
    target: Optional[JobTarget] = None,
    use_enhanced: bool = True,
):
    """
    Genera un documento Word (.docx) estrictamente lineal, sin tablas,
    sin cuadros de texto, usando fuentes estándar ATS-friendly.

    Args:
        cv: Datos estructurados del CV (CVData).
        output_path: Ruta donde guardar el archivo .docx.
        target: Puesto objetivo (opcional, para metadata).
        use_enhanced: Si True, usa campos mejorados por IA cuando estén disponibles.

    Raises:
        IOError: Si no se puede escribir en la ruta de salida.
    """
    doc = DocumentBuilder()

    # --- Configurar márgenes estilo carta ---
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

# --- Estilo base del párrafo ---
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME  # type: ignore[attr-defined]
    style.font.size = Pt(FONT_SIZE_BODY)  # type: ignore[attr-defined]
    style.paragraph_format.space_after = Pt(4)  # type: ignore[attr-defined]
    style.paragraph_format.space_before = Pt(0)  # type: ignore[attr-defined]
    style.paragraph_format.line_spacing = 1.15  # type: ignore[attr-defined]

    # ================================================================
    # 1. DATOS DE CONTACTO (Nombre + información)
    # ================================================================
    lines = cv.datos_contacto.split("\n")
    name_line = lines[0] if lines else "Nombre no especificado"
    contact_rest = "\n".join(lines[1:]) if len(lines) > 1 else ""

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_name.add_run(name_line.upper())
    _set_font(run, size=FONT_SIZE_TITLE, bold=True, color=COLOR_DARK)

    if contact_rest:
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_contact.add_run(contact_rest.strip())
        _set_font(run, size=10, color=RGBColor(0x44, 0x44, 0x44))

    # Metadata si hay puesto objetivo
    if target:
        p_target = doc.add_paragraph()
        p_target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_target.add_run(f"Puesto Objetivo: {target.titulo}")
        _set_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))

    _add_separator(doc)

    # ================================================================
    # 2. PERFIL PROFESIONAL
    # ================================================================
    _add_section_title(doc, "Perfil Profesional")
    perfil = (cv.perfil_mejorado if (use_enhanced and cv.perfil_mejorado)
              else cv.perfil)
    p_perfil = doc.add_paragraph()
    run = p_perfil.add_run(perfil)
    _set_font(run, size=FONT_SIZE_BODY)
    _add_separator(doc)

    # ================================================================
    # 3. EXPERIENCIA LABORAL
    # ================================================================
    _add_section_title(doc, "Experiencia Laboral")

    for exp in cv.experiencia_laboral:
        # Encabezado del puesto
        p_head = doc.add_paragraph()
        run = p_head.add_run(f"{exp.puesto}")
        _set_font(run, size=FONT_SIZE_BODY, bold=True)

        # Empresa y fechas
        p_meta = doc.add_paragraph()
        run = p_meta.add_run(f"{exp.empresa}  |  {exp.fechas}")
        _set_font(run, size=10, color=RGBColor(0x55, 0x55, 0x55))

        # Logros (mejorados si están disponibles)
        logros = exp.logros_mejorados if (use_enhanced and exp.logros_mejorados) else exp.logros
        for logro in logros:
            p_bullet = doc.add_paragraph(style="List Bullet")
            run = p_bullet.add_run(logro)
            _set_font(run, size=FONT_SIZE_BODY)

    _add_separator(doc)

    # ================================================================
    # 4. EDUCACIÓN
    # ================================================================
    _add_section_title(doc, "Educación")

    for edu in cv.educacion:
        p_head = doc.add_paragraph()
        run = p_head.add_run(edu.titulo)
        _set_font(run, size=FONT_SIZE_BODY, bold=True)

        p_meta = doc.add_paragraph()
        run = p_meta.add_run(f"{edu.institucion}  |  {edu.fechas}")
        _set_font(run, size=10, color=RGBColor(0x55, 0x55, 0x55))

    _add_separator(doc)

    # ================================================================
    # 5. HABILIDADES
    # ================================================================
    _add_section_title(doc, "Habilidades")

    # Palabras clave ATS
    if use_enhanced and cv.palabras_clave_ats:
        p_keywords = doc.add_paragraph()
        run = p_keywords.add_run(
            "Keywords ATS: " + ", ".join(cv.palabras_clave_ats)
        )
        _set_font(run, size=9, color=RGBColor(0x55, 0x55, 0x55))
        doc.add_paragraph()  # espacio

    habilidades = cv.habilidades_mejoradas if (use_enhanced and cv.habilidades_mejoradas) else cv.habilidades
    for skill in habilidades:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(skill)
        _set_font(run, size=FONT_SIZE_BODY)

    # --- Guardar el documento ---
    output_path_obj = Path(output_path)
    output_path_obj.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path_obj))
    log.info("Documento .docx guardado en: %s", output_path)


# ============================================================================
# GENERADOR .HTML (Template ATS)
# ============================================================================

HTML_BASE_TEMPLATE = r"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <meta name="description" content="CV ATS - {NOMBRE} - {PUESTO}">
  <meta property="og:title" content="CV - {NOMBRE}">
  <meta property="og:description" content="Currículum Vitae profesional optimizado para ATS">
  <meta property="og:type" content="profile">
  <title>CV - {NOMBRE} | {PUESTO}</title>
  <script type="application/ld+json">
  {{
    "@context": "https://schema.org",
    "@type": "Person",
    "name": "{NOMBRE}",
    "jobTitle": "{PUESTO}",
    "description": "{PERFIL_JSON}",
    "knowsAbout": {HABILIDADES_JSON}
  }}
  </script>
  <style>
    :root {{
      --sidebar-width: 32%;
      --sidebar-bg: #1a232a;
      --sidebar-main-bg: #ffffff;
      --accent-color: #0d6efd;
      --text-dark: #111111;
      --text-light: #e0e6ed;
    }}

    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
      font-family: 'Segoe UI', system-ui, -apple-system, sans-serif;
    }}

    body {{
      background-color: #eef2f5;
      color: #333;
      display: flex;
      justify-content: center;
      padding: 20px;
    }}

    .btn-export {{
      position: fixed;
      top: 20px;
      right: 20px;
      background-color: var(--accent-color);
      color: #ffffff;
      border: none;
      padding: 12px 20px;
      border-radius: 6px;
      font-weight: 600;
      font-size: 14px;
      cursor: pointer;
      box-shadow: 0 4px 12px rgba(0,0,0,0.15);
      transition: background-color 0.2s ease, transform 0.1s ease;
      z-index: 1000;
    }}

    .btn-export:hover {{
      background-color: #0b5ed7;
      transform: translateY(-1px);
    }}

    .cv-container {{
      width: 210mm;
      min-height: 297mm;
      background: #ffffff;
      display: flex;
      box-shadow: 0 8px 24px rgba(0, 0, 0, 0.1);
    }}

    .sidebar {{
      width: var(--sidebar-width, 32%);
      background-color: var(--sidebar-bg, #1a232a);
      color: var(--text-light);
      padding: 35px 20px;
      display: flex;
      flex-direction: column;
      gap: 25px;
    }}

    .sidebar-section {{
      page-break-inside: avoid;
      break-inside: avoid;
    }}

    .sidebar-section h3 {{
      font-size: 12px;
      letter-spacing: 1.2px;
      text-transform: uppercase;
      color: #ffffff;
      margin-bottom: 12px;
      border-bottom: 1px solid #344454;
      padding-bottom: 5px;
      font-weight: 700;
      page-break-after: avoid;
      break-after: avoid;
    }}

    .sidebar-section p {{
      font-size: 10.5px;
      line-height: 1.5;
      color: #c2cbd6;
      text-align: justify;
    }}

    .skills-list, .sidebar-list {{
      list-style: none;
    }}

    .skills-list li, .sidebar-list li {{
      font-size: 10px;
      padding: 3.5px 0;
      border-bottom: 1px solid #232e3a;
      text-transform: uppercase;
      letter-spacing: 0.5px;
      color: #d1d9e2;
    }}

    .sidebar-list li strong {{
      color: #ffffff;
      display: block;
      margin-bottom: 2px;
    }}

    .main-content {{
      width: calc(100% - var(--sidebar-width, 32%));
      padding: 40px 30px;
      background-color: var(--sidebar-main-bg, #ffffff);
    }}

    .header {{
      border-bottom: 2px solid var(--text-dark);
      padding-bottom: 15px;
      margin-bottom: 25px;
    }}

    .header h1 {{
      font-size: 24px;
      font-weight: 800;
      letter-spacing: 1px;
      color: var(--text-dark);
      line-height: 1.1;
      text-transform: uppercase;
      margin-bottom: 8px;
    }}

    .contact-info {{
      font-size: 10.5px;
      color: #444;
      line-height: 1.5;
    }}

    .job-target-badge {{
      display: inline-block;
      background-color: var(--accent-color);
      color: white;
      padding: 4px 10px;
      border-radius: 4px;
      font-size: 10px;
      margin-top: 6px;
      font-weight: 600;
    }}

    .section-title {{
      font-size: 13px;
      font-weight: 800;
      letter-spacing: 1.5px;
      text-transform: uppercase;
      color: var(--text-dark);
      border-bottom: 1.5px solid var(--text-dark);
      padding-bottom: 4px;
      margin-bottom: 18px;
      margin-top: 15px;
    }}

    .timeline-item {{
      display: flex;
      margin-bottom: 16px;
      page-break-inside: avoid;
      break-inside: avoid;
    }}

    .company-col {{
      width: 34%;
      padding-right: 10px;
    }}

    .company-name {{
      font-size: 11px;
      font-weight: 800;
      color: var(--text-dark);
      text-transform: uppercase;
    }}

    .company-location, .item-date {{
      font-size: 10px;
      color: #666;
    }}

    .details-col {{
      width: 66%;
      padding-left: 15px;
      border-left: 1.5px solid #e0e0e0;
      position: relative;
    }}

    .details-col::before {{
      content: '';
      position: absolute;
      left: -4.5px;
      top: 4px;
      width: 7px;
      height: 7px;
      background-color: var(--sidebar-bg);
      border-radius: 50%;
    }}

    .role-title {{
      font-size: 11.5px;
      font-weight: 700;
      color: var(--text-dark);
      margin-bottom: 5px;
    }}

    .role-bullets {{
      list-style-type: disc;
      padding-left: 14px;
    }}

    .role-bullets li {{
      font-size: 10px;
      line-height: 1.4;
      color: #444;
      margin-bottom: 3px;
    }}

    .keywords-ats {{
      display: flex;
      flex-wrap: wrap;
      gap: 4px;
      margin-top: 10px;
    }}

    .keywords-ats span {{
      background-color: #e8f0fe;
      color: #1a232a;
      font-size: 8.5px;
      padding: 3px 8px;
      border-radius: 10px;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }}

    @media print {{
      body {{
        background: none !important;
        padding: 0 !important;
        -webkit-print-color-adjust: exact;
        print-color-adjust: exact;
        color-adjust: exact;
      }}
      .btn-export {{
        display: none !important;
      }}
      .no-print {{
        display: none !important;
      }}
      body::before {{
        content: '';
        position: fixed;
        top: 0;
        left: 0;
        width: var(--sidebar-width, 32%);
        height: 100%;
        background-color: var(--sidebar-bg, #1a232a) !important;
        z-index: -1;
        -webkit-print-color-adjust: exact;
        print-color-adjust: exact;
        color-adjust: exact;
      }}
      body::after {{
        content: 'Curriculum Vitae - {NOMBRE}';
        position: fixed;
        bottom: 5mm;
        left: 0;
        width: 100%;
        text-align: center;
        font-size: 8px;
        font-family: Arial, sans-serif;
        color: #888888;
        z-index: 0;
        pointer-events: none;
        -webkit-print-color-adjust: exact;
        print-color-adjust: exact;
        color-adjust: exact;
      }}
      .cv-container {{
        box-shadow: none !important;
        width: 100% !important;
        min-height: 100vh !important;
        page-break-after: avoid;
        background: transparent !important;
      }}
      .sidebar {{
        background: transparent !important;
        color: var(--text-light) !important;
      }}
      .main-content {{
        background: transparent !important;
      }}
      .keywords-ats span {{
        background-color: #2a3a4a !important;
        color: #d1d9e2 !important;
      }}
      @page {{
        margin-top: 15mm;
        margin-bottom: 15mm;
        margin-left: 0;
        margin-right: 0;
        size: A4;
      }}
    }}

    @media screen and (max-width: 768px) {{
      body {{
        padding: 10px;
      }}
      .cv-container {{
        flex-direction: column;
        width: 100%;
        min-height: auto;
      }}
      .sidebar {{
        width: 100%;
        padding: 20px;
      }}
      .main-content {{
        width: 100%;
        padding: 20px;
      }}
      .timeline-item {{
        flex-direction: column;
      }}
      .company-col, .details-col {{
        width: 100%;
      }}
      .details-col {{
        padding-left: 0;
        border-left: none;
        margin-top: 5px;
      }}
      .details-col::before {{
        display: none;
      }}
      .btn-export {{
        top: 10px;
        right: 10px;
        padding: 8px 14px;
        font-size: 12px;
      }}
    }}
  </style>
</head>
<body>
  <button class="btn-export" onclick="window.print()" aria-label="Exportar a PDF">
    Exportar a PDF
  </button>

  <main class="cv-container" itemscope itemtype="https://schema.org/Person">
    
    <aside class="sidebar" role="complementary" aria-label="Información complementaria">
      <section class="sidebar-section">
        <h3>Sobre Mí</h3>
        <p itemprop="description">{PERFIL}</p>
      </section>

      <section class="sidebar-section">
        <h3>Habilidades</h3>
        <ul class="skills-list" itemprop="knowsAbout">
          {HABILIDADES}
        </ul>
      </section>

      <section class="sidebar-section">
        <h3>Idiomas</h3>
        <ul class="sidebar-list">
          <li><strong>Inglés:</strong> Nivel Avanzado</li>
        </ul>
      </section>

      {SIDEBAR_EXTRA}

      <section class="sidebar-section">
        <h3>Permiso de Conducir</h3>
        <ul class="sidebar-list">
          <li>Licencia Clase B vigente</li>
        </ul>
      </section>
    </aside>

    <article class="main-content" role="main" aria-label="Contenido principal del CV">
      <header class="header">
        <h1 itemprop="name">{NOMBRE_HTML}</h1>
        <div class="contact-info" itemprop="address">
          {CONTACTO}
        </div>
        <span class="job-target-badge" aria-label="Puesto objetivo">
          🎯 {PUESTO}
        </span>
      </header>

      <section aria-label="Experiencia Laboral">
        <h2 class="section-title">Experiencia Laboral</h2>
        {EXPERIENCIA}
      </section>

      <section aria-label="Educación">
        <h2 class="section-title">Educación</h2>
        {EDUCACION}
      </section>

      {KEYWORDS_SECTION}
    </article>
  </main>
</body>
</html>"""


def generate_ats_html(
    cv: CVData,
    output_path: str,
    target: Optional[JobTarget] = None,
    use_enhanced: bool = True,
):
    """
    Genera una página HTML del CV con diseño ATS, semántica HTML5,
    microdatos Schema.org y diseño responsive.

    Args:
        cv: Datos estructurados del CV.
        output_path: Ruta donde guardar el archivo .html.
        target: Puesto objetivo (opcional).
        use_enhanced: Si True, usa campos mejorados por IA.
    """
    # --- Parsear nombre ---
    nombre_raw = cv.datos_contacto.split("\n")[0] if cv.datos_contacto else "Nombre Completo"
    nombre_parts = nombre_raw.strip().split()
    nombre_html = "<br>".join(
        [" ".join(nombre_parts[:2]), " ".join(nombre_parts[2:])]
    ) if len(nombre_parts) > 2 else nombre_raw

    # --- Perfil ---
    perfil = (cv.perfil_mejorado if (use_enhanced and cv.perfil_mejorado) else cv.perfil)
    perfil_escaped = perfil[:150].replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "\\r") if perfil else ""

    # --- Contacto ---
    contacto_lines = cv.datos_contacto.split("\n")
    contacto = "<br>".join(contacto_lines[1:]) if len(contacto_lines) > 1 else ""

    # --- Habilidades ---
    habilidades = cv.habilidades_mejoradas if (use_enhanced and cv.habilidades_mejoradas) else cv.habilidades
    habilidades_html = "\n          ".join(
        f"<li>{h}</li>" for h in habilidades
    )
    habilidades_json = json.dumps(habilidades)

    # --- Sidebar extra (prácticas) ---
    sidebar_extra = ""
    for exp in cv.experiencia_laboral:
        if "practica" in exp.puesto.lower() or "practicante" in exp.puesto.lower():
            logros_text = "<br>".join(exp.logros_mejorados if (use_enhanced and exp.logros_mejorados) else exp.logros)
            sidebar_extra += f"""
      <section class="sidebar-section">
        <h3>Prácticas</h3>
        <ul class="sidebar-list">
          <li>
            <strong>{exp.empresa.upper()}</strong>
            <span style="font-size: 9px; color: #a0aec0;">{exp.fechas}</span>
            <p style="margin-top: 4px;">{logros_text}</p>
          </li>
        </ul>
      </section>"""

    # --- Experiencia laboral ---
    experiencia_html = ""
    for exp in cv.experiencia_laboral:
        if "practica" in exp.puesto.lower() or "practicante" in exp.puesto.lower():
            continue
        logros = exp.logros_mejorados if (use_enhanced and exp.logros_mejorados) else exp.logros
        bullets = "\n              ".join(
            f"<li>{l}</li>" for l in logros
        )
        experiencia_html += f"""
        <div class="timeline-item">
          <div class="company-col">
            <div class="company-name" itemprop="name">{exp.empresa}</div>
            <div class="item-date">{exp.fechas}</div>
          </div>
          <div class="details-col">
            <div class="role-title">{exp.puesto}</div>
            <ul class="role-bullets">
              {bullets}
            </ul>
          </div>
        </div>"""

    # --- Educación ---
    educacion_html = ""
    for edu in cv.educacion:
        educacion_html += f"""
        <div class="timeline-item">
          <div class="company-col">
            <div class="company-name">{edu.institucion}</div>
            <div class="item-date">{edu.fechas}</div>
          </div>
          <div class="details-col">
            <div class="role-title">{edu.titulo}</div>
          </div>
        </div>"""

    # --- Keywords ATS ---
    keywords_html = ""
    if use_enhanced and cv.palabras_clave_ats:
        kw_badges = " ".join(f'<span>{kw}</span>' for kw in cv.palabras_clave_ats)
        keywords_html = f"""
      <section aria-label="Palabras clave ATS">
        <h2 class="section-title">Keywords ATS</h2>
        <div class="keywords-ats">
          {kw_badges}
        </div>
      </section>"""

    # --- Puesto objetivo ---
    puesto = target.titulo if target else "No especificado"

    # --- Renderizar template ---
    html_content = (
        HTML_BASE_TEMPLATE
        .replace("{NOMBRE}", nombre_raw)
        .replace("{NOMBRE_HTML}", nombre_html)
        .replace("{PERFIL}", perfil)
        .replace("{PERFIL_JSON}", perfil_escaped)
        .replace("{CONTACTO}", contacto)
        .replace("{HABILIDADES}", habilidades_html)
        .replace("{HABILIDADES_JSON}", habilidades_json)
        .replace("{PUESTO}", puesto)
        .replace("{SIDEBAR_EXTRA}", sidebar_extra)
        .replace("{EXPERIENCIA}", experiencia_html)
        .replace("{EDUCACION}", educacion_html)
        .replace("{KEYWORDS_SECTION}", keywords_html)
    )

    # --- Guardar ---
    output_path_obj = Path(output_path)
    output_path_obj.parent.mkdir(parents=True, exist_ok=True)
    output_path_obj.write_text(html_content, encoding="utf-8")
    log.info("Documento HTML guardado en: %s", output_path)


def _add_section_title(doc, title: str):
    """Agrega un título de sección en mayúsculas con formato consistente."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title.upper())
    _set_font(run, size=FONT_SIZE_SECTION, bold=True, color=COLOR_DARK)


def _add_separator(doc):
    """Agrega una línea separadora horizontal sutil."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("─" * 65)
    _set_font(run, size=8, color=RGBColor(0xCC, 0xCC, 0xCC))




